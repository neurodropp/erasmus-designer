#!/usr/bin/env python3
"""Read user XLSX files without modifying/recalculating them; inventory all tabs/cells.

Requires bundled Python/openpyxl. Cached values are historical snapshots, not proof of
recalculation. ZIP XML preserves comments/merges/visibility unavailable in read-only mode.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import posixpath
import shutil
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}
FILES = [('template-task-file-2024', 'TASK FILE - 2024 TEMPLATE.xlsx', 'task-file-2024', 2024),
         ('template-deor-budget-user', 'DEOR BY DESIGN BUDGET.xlsx', 'deor-budget', None)]
ERRORS = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!'}

def val(v):
    if isinstance(v, (datetime, date)): return v.isoformat()
    if v is None or isinstance(v, (str, int, float, bool)): return v
    return str(v)

def text(v):
    return str(v).replace('|', '\\|').replace('\r', '').replace('\n', '<br>')

def xml_meta(z, sheet_path):
    root = ET.fromstring(z.read(sheet_path))
    relpath = posixpath.join(posixpath.dirname(sheet_path), '_rels', posixpath.basename(sheet_path)+'.rels')
    comments, rels = {}, []
    if relpath in z.namelist():
        for rel in ET.fromstring(z.read(relpath)):
            rels.append(dict(rel.attrib))
            if rel.attrib.get('Type', '').endswith('/comments'):
                target = rel.attrib['Target']
                cp = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), target))
                cr = ET.fromstring(z.read(cp))
                authors = [n.text or '' for n in cr.findall('s:authors/s:author', NS)]
                for c in cr.findall('s:commentList/s:comment', NS):
                    comments[c.attrib['ref']] = {'author': authors[int(c.attrib.get('authorId',0))],
                      'text': ''.join(n.text or '' for n in c.findall('.//s:t', NS))}
    for hyperlink in root.findall('s:hyperlinks/s:hyperlink', NS):
        rid=hyperlink.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        match=next((r for r in rels if r.get('Id')==rid),None)
        if match: hyperlink.attrib['resolved_target']=match.get('Target','')
    return {
        'comments': comments, 'relationships': rels,
        'merged_ranges': [m.attrib['ref'] for m in root.findall('s:mergeCells/s:mergeCell', NS)],
        'hidden_rows': [int(r.attrib['r']) for r in root.findall('s:sheetData/s:row', NS) if r.attrib.get('hidden')=='1'],
        'hidden_columns': [dict(c.attrib) for c in root.findall('s:cols/s:col', NS) if c.attrib.get('hidden')=='1'],
        'validations': [dict(d.attrib, formulas=[n.text for n in d]) for d in root.findall('s:dataValidations/s:dataValidation', NS)],
        'conditional_formats': [dict(c.attrib) for c in root.findall('s:conditionalFormatting', NS)],
        'hyperlinks': [dict(h.attrib) for h in root.findall('s:hyperlinks/s:hyperlink', NS)],
        'print_options': [dict(p.attrib) for tag in ['pageSetup','pageMargins','printOptions'] for p in root.findall('s:'+tag, NS)],
        'threaded_comment_parts': [r for r in rels if 'threadedComment' in r.get('Type','')],
    }

def extract(source, dest, source_id, slug, publication_year):
    dest.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    copy = dest / source.name
    shutil.copy2(source, copy)
    w = load_workbook(source, read_only=True, data_only=False)
    cached = load_workbook(source, read_only=True, data_only=True)
    result = {'source_id':source_id,'original_path':str(source),'archived_path':str(copy),
      'sha256':digest,'official':False,'publication_year':publication_year,
      'financial_rates_verified':False,'recalculated':False,'sheets':[],
      'defined_names':{k:{'formula':v.attr_text,'scope_sheet_index':v.localSheetId,'hidden':v.hidden} for k,v in w.defined_names.items()},
      'calculation_properties':str(w.calculation),'warnings':[]}
    dump = [f'# {source.name}', '', 'User-provided, non-normative source. All values are examples, template content, or historical data not confirmed for a new project. No rates have been verified for 2026. Formulas are preserved and have not been recalculated. The original workbook was not modified.', '',f'SHA256: `{digest}`', '']
    with ZipFile(source) as z:
        for idx, ws in enumerate(w,1):
            path=ws._worksheet_path
            meta=xml_meta(z,path)
            cws=cached[ws.title]
            # Missing dimensions occur in these supplied XLSX; streaming rows discovers actual cells.
            rows=list(ws.iter_rows())
            crs=list(cws.iter_rows())
            cache_map={c.coordinate:c.value for row in crs for c in row if hasattr(c,'coordinate')}
            cells=[]
            for row in rows:
                for c in row:
                    if c.value is None: continue
                    formula_value = c.value.text if c.data_type=='f' and hasattr(c.value,'text') else c.value
                    record={'cell':c.coordinate,'value':val(formula_value),'type':c.data_type,
                      'array_ref':getattr(c.value,'ref',None) if c.data_type=='f' else None,
                      'cached_value':val(cache_map.get(c.coordinate)) if c.data_type=='f' else None,
                      'number_format':c.number_format,'bold':bool(c.font.bold),
                      'font_color':str(c.font.color),'fill':str(c.fill.fgColor),
                      'comment':meta['comments'].get(c.coordinate)}
                    cells.append(record)
            coords={c['cell'] for c in cells}
            for coord, comment in meta['comments'].items():
                if coord not in coords: cells.append({'cell':coord,'value':None,'type':'comment-only','comment':comment,'cached_value':None})
            occupied=[c for row in rows for c in row if getattr(c,'value',None) is not None]
            bounds=(f'{get_column_letter(min(c.column for c in occupied))}{min(c.row for c in occupied)}:'
                    f'{get_column_letter(max(c.column for c in occupied))}{max(c.row for c in occupied)}') if occupied else None
            formulas=[c for c in cells if c['type']=='f']
            errors=[c for c in cells if c['type']=='e' or c['value'] in ERRORS or c.get('cached_value') in ERRORS]
            no_cache=[c['cell'] for c in formulas if c['cached_value'] is None]
            headings=[{'cell':c['cell'],'text':c['value']} for c in cells if isinstance(c['value'],str) and c['type']!='f' and c.get('bold')]
            item={'index':idx,'name':ws.title,'state':ws.sheet_state,'occupied_range':bounds,
              'populated_cells':len(occupied),'formulas_count':len(formulas),'comments_count':len(meta['comments']),
              'known_errors_count':len(errors),'errors':errors,'formula_missing_cache':no_cache,
              'headings_candidates':headings,'cells':cells, **meta}
            result['sheets'].append(item)
            out=Path('sources/templates-workbooks')/slug
            out.mkdir(parents=True,exist_ok=True)
            sheet_slug=f'{idx:02d}'
            (out/(sheet_slug+'.json')).write_text(json.dumps(item,ensure_ascii=False,indent=2)+'\n')
            sd=[f'## {idx}. {ws.title}', '',f'State: {ws.sheet_state}; occupied range: {bounds}; populated cells: {len(occupied)}; formulas: {len(formulas)}; comments: {len(meta["comments"])}; known stored errors: {len(errors)}.', '',
                f'Merged ranges: {", ".join(meta["merged_ranges"]) or "none"}. Hidden rows: {meta["hidden_rows"]}. Hidden columns: {meta["hidden_columns"]}.', '',
                '| Cell | Exact content / formula | Stored formula value | Comment |','|---|---|---|---|']
            for c in cells:
                comment=c.get('comment')
                sd.append(f'| {c["cell"]} | {text(c["value"])} | {text(c["cached_value"]) if c["type"]=="f" else ""} | {text(comment["author"]+": "+comment["text"]) if comment else ""} |')
            sd.extend(['', '### Sheet links, validation, and metadata', '',
              '```json', json.dumps({k:item[k] for k in ['hyperlinks','validations','conditional_formats','print_options','hidden_rows','hidden_columns']},ensure_ascii=False,indent=2), '```'])
            (out/(sheet_slug+'.md')).write_text('\n'.join(sd)+'\n')
            dump.extend(sd+[''])
    dump.extend(['## Workbook named ranges and functions', '', '| Name | Exact definition |', '|---|---|'])
    for name, definition in result['defined_names'].items():
        dump.append(f'| {text(name)} | {text(definition["formula"])} |')
    w.close(); cached.close()
    result['totals']={k:sum(s[k] for s in result['sheets']) for k in ['populated_cells','formulas_count','comments_count','known_errors_count']}
    result['totals']['sheet_count']=len(result['sheets'])
    result['totals']['hidden_sheets_count']=sum(s['state']!='visible' for s in result['sheets'])
    result['totals']['formula_missing_cache_count']=sum(len(s['formula_missing_cache']) for s in result['sheets'])
    result['totals']['array_formula_count']=sum(bool(c.get('array_ref')) for s in result['sheets'] for c in s['cells'])
    result['totals']['dummyfunction_formula_count']=sum('__xludf.DUMMYFUNCTION' in str(c['value']) for s in result['sheets'] for c in s['cells'] if c['type']=='f')
    result['totals']['threaded_comment_parts_count']=sum(len(s['threaded_comment_parts']) for s in result['sheets'])
    (dest/(slug+'-extracted.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    (dest/(slug+'-extracted.md')).write_text('\n'.join(dump)+'\n')
    assert hashlib.sha256(source.read_bytes()).hexdigest()==digest, 'Source changed unexpectedly'
    assert hashlib.sha256(copy.read_bytes()).hexdigest()==digest, 'Archive differs from the source'
    return result

def main():
    ap=argparse.ArgumentParser(description='Extract workbook sheets, cells, formulas, and metadata without modifying the source workbooks.')
    ap.add_argument('--input-dir',type=Path,required=True,help='Directory containing the two original workbook files; extraction rewrites generated outputs.')
    ap.add_argument('--archive-dir',type=Path,default=Path('raw/templates'),help='Directory for archived copies of the original workbooks (default: raw/templates).')
    args=ap.parse_args()
    records=[]
    for source_id,name,slug,year in FILES:
        r=extract(args.input_dir/name,args.archive_dir,source_id,slug,year)
        records.append({k:r[k] for k in ['source_id','sha256','original_path','archived_path','totals']})
        print(f'Extracted {source_id}: {json.dumps(r["totals"])}')
    Path('sources/templates-workbooks').mkdir(parents=True,exist_ok=True)
    Path('sources/templates-workbooks/manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')

if __name__=='__main__': main()
